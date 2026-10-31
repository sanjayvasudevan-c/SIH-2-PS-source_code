"""Import flow: Demo A counts, guided corrections, exclusion, commit, retries and file checks."""

import io

from openpyxl import Workbook
from sqlalchemy import update

from app.db import SessionLocal
from app.models import ImportJob

from .conftest import SAMPLES, rows_by_roll, upload

DEMO_A = SAMPLES / "demo_A_kaveri_upload_2026-27.csv"
DEMO_A_XLSX = SAMPLES / "demo_A_kaveri_upload_2026-27.xlsx"
STAFF = "kaveri.staff@demo.local"


def counts(job):
    c = job["summary"]["counts"]
    return c["VALID"], c["REVIEW"], c["ERROR"]


def test_demo_a_preview_counts_csv_and_xlsx(client, as_):
    h = as_(STAFF)
    r = upload(client, h, DEMO_A)
    assert r.status_code == 200, r.text
    assert counts(r.json()) == (5, 2, 9)
    rx = upload(client, h, DEMO_A_XLSX)
    assert rx.status_code == 200, rx.text
    assert counts(rx.json()) == (5, 2, 9)
    # same CSV again -> file-level F004 warning, row counts unchanged
    again = upload(client, h, DEMO_A).json()
    assert counts(again) == (5, 2, 9)
    assert [f["rule_id"] for f in again["summary"]["file_findings"]] == ["F004_REPEATED_FILE"]
    for j in (r.json(), rx.json(), again):
        client.post(f"/api/v1/imports/{j['id']}/discard", headers=h)


def test_demo_a_full_flow(client, as_):
    h = as_(STAFF)
    job = upload(client, h, DEMO_A).json()
    jid = job["id"]
    rows = rows_by_roll(client, h, jid)
    # did-you-mean for Sanjana
    sanjana = rows["KIT26ECE908"]
    sugg = next(f["suggestion"] for f in sanjana["findings"] if f.get("suggestion"))
    assert sugg["value"] == "911100000008"
    fixes = {
        "KIT26ECE908": {"apaar_id": sugg["value"]},
        "KIT26CSE906": {"apaar_id": "911100000006"},
        "KIT26ECE907": {"apaar_id": "911100000007"},
        "KIT26ECE910": {"date_of_birth": "2008-06-18"},
        "KIT26CSE913": {"seat_category": "REGULAR"},
        "KIT26MTA914": {"programme_code": "MTECH-CSE"},
        "KIT26ECE915": {"gender": "F"},
    }
    for roll, ch in fixes.items():
        r = client.patch(f"/api/v1/imports/{jid}/rows/{rows[roll]['row_no']}", headers=h, json={"changes": ch})
        assert r.status_code == 200, r.text
        assert r.json()["row"]["status"] in ("VALID", "WARNING", "REVIEW", "ERROR")
    rows = rows_by_roll(client, h, jid)
    for roll in fixes:
        assert rows[roll]["status"] == "VALID", (roll, rows[roll]["findings"])
        assert rows[roll]["corrected"] is True and rows[roll]["original"] != rows[roll]["data"]
    # leave out the duplicate Riya row -> the other Riya row becomes valid
    r = client.post(f"/api/v1/imports/{jid}/rows/{rows['KIT26CSE912']['row_no']}/exclude", headers=h,
                    json={"excluded": True})
    assert r.status_code == 200
    rows = rows_by_roll(client, h, jid)
    assert rows["KIT26CSE912"]["status"] == "EXCLUDED" and rows["KIT26CSE911"]["status"] == "VALID"
    job = client.get(f"/api/v1/imports/{jid}", headers=h).json()
    assert job["summary"]["committable"] == 15 and job["summary"]["counts"]["ERROR"] == 0
    proj = {p["programme_code"]: p for p in job["summary"]["projection"]}
    assert proj["BTECH-ECE"]["after"]["within_intake_enrolled"] == 60 and not proj["BTECH-ECE"]["exceeds"]
    # commit 15
    r = client.post(f"/api/v1/imports/{jid}/commit", headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["committed_now"] == 15 and r.json()["status"] == "COMMITTED"
    # re-commit is a no-op
    r2 = client.post(f"/api/v1/imports/{jid}/commit", headers=h)
    assert r2.status_code == 200 and r2.json()["committed_now"] == 0
    # committed row not editable
    r3 = client.patch(f"/api/v1/imports/{jid}/rows/{rows['KIT26CSE901']['row_no']}", headers=h,
                      json={"changes": {"gender": "M"}})
    assert r3.status_code == 409
    assert r3.json()["error"]["status"] == 409
    # R007 (Vignesh) and R004 (Tomas) flags opened for AICTE
    flags = client.get("/api/v1/issues?status=OPEN", headers=as_("aicte.admin@demo.local")).json()["items"]
    rules = {f["rule_id"] for f in flags if f["institution_name"] and "Kaveri" in f["institution_name"]}
    assert {"R007_APAAR_MISMATCH", "R004_APAAR_PENDING"} <= rules
    # check report CSV
    rep = client.get(f"/api/v1/imports/{jid}/report.csv", headers=h)
    assert rep.status_code == 200 and "spreadsheet_row" in rep.text


def test_partial_commit_then_retry(client, as_):
    h = as_("brahmaputra.staff@demo.local")
    csv = ("institution_student_id,full_name,date_of_birth,gender,nationality_status,programme_code,academic_year,"
           "seat_category,admission_date,apaar_id,apaar_pending_reason\n"
           "BPE26P1,Partha Ghosh,2008-01-05,M,INDIAN,BTECH-CIVIL,2026-27,REGULAR,2026-08-10,,APPLIED_AWAITING_ID\n"
           "BPE26P2,Rina Ghosh,2008-01-06,F,INDIAN,BTECH-CIVIL,2026-27,REGULAR,2026-08-10,,\n")
    job = upload(client, h, None, name="partial.csv", content=csv.encode()).json()
    assert counts(job) == (0, 1, 1)
    r = client.post(f"/api/v1/imports/{job['id']}/commit", headers=h).json()
    assert r["committed_now"] == 1 and r["status"] == "PARTIALLY_COMMITTED"
    rows = rows_by_roll(client, h, job["id"])
    # a committed row inside a partially committed job cannot be edited (row-level 409)
    r409 = client.patch(f"/api/v1/imports/{job['id']}/rows/{rows['BPE26P1']['row_no']}", headers=h,
                        json={"changes": {"gender": "F"}})
    assert r409.status_code == 409
    assert r409.json()["error"]["status"] == 409 and r409.json()["error"]["row_status"] == "COMMITTED"
    client.patch(f"/api/v1/imports/{job['id']}/rows/{rows['BPE26P2']['row_no']}", headers=h,
                 json={"changes": {"apaar_pending_reason": "APPLIED_AWAITING_ID"}})
    r = client.post(f"/api/v1/imports/{job['id']}/commit", headers=h).json()
    assert r["committed_now"] == 1 and r["status"] == "COMMITTED" and r["committed_total"] == 2


def test_concurrent_commit_conflict(client, as_):
    h = as_(STAFF)
    job = upload(client, h, DEMO_A).json()
    with SessionLocal() as s:
        s.execute(update(ImportJob).where(ImportJob.id == job["id"]).values(status="COMMITTING"))
        s.commit()
    r = client.post(f"/api/v1/imports/{job['id']}/commit", headers=h)
    assert r.status_code == 409
    with SessionLocal() as s:
        s.execute(update(ImportJob).where(ImportJob.id == job["id"]).values(status="PREVIEW"))
        s.commit()
    client.post(f"/api/v1/imports/{job['id']}/discard", headers=h)


def test_bulk_fix(client, as_):
    h = as_(STAFF)
    job = upload(client, h, DEMO_A).json()
    r = client.post(f"/api/v1/imports/{job['id']}/bulk-fix", headers=h,
                    json={"field": "programme_code", "from": "MTECH-AI", "to": "MTECH-CSE"})
    assert r.status_code == 200 and r.json()["updated_rows"] == 1
    assert r.json()["job"]["summary"]["counts"]["ERROR"] == 8
    client.post(f"/api/v1/imports/{job['id']}/discard", headers=h)


def test_upload_rejections(client, as_):
    h = as_(STAFF)
    assert upload(client, h, None, name="data.txt", content=b"a,b").status_code == 415
    assert upload(client, h, None, name="fake.xlsx", content=b"not a zip").status_code == 415
    assert upload(client, h, None, name="nul.csv", content=b"a,b\x00\n").status_code == 415
    big = b"full_name\n" + b"x" * (2 * 1024 * 1024 + 10)
    assert upload(client, h, None, name="big.csv", content=big).status_code == 413
    r = upload(client, h, None, name="bad.csv", content=b"name,dob\nA,2008-01-01\n")
    assert r.status_code == 422 and "missing required columns" in r.json()["error"]["message"]
    # header aliases accepted
    ok = (b"apaar,roll_no,name,dob,gender,nationality,programme,year,seat,doa\n"
          b"911100000001,KIT26AL1,Meenakshi Sundaram,12/05/2008,F,Indian,BTECH-CSE,2026-27,REGULAR,03-08-2026\n")
    r = upload(client, h, None, name="alias.csv", content=ok)
    assert r.status_code == 200, r.text
    client.post(f"/api/v1/imports/{r.json()['id']}/discard", headers=h)


def test_template_downloads(client, as_):
    h = as_(STAFF)
    r = client.get("/api/v1/imports/template.csv", headers=h)
    assert r.status_code == 200 and r.text.startswith("apaar_id,institution_student_id")
    r = client.get("/api/v1/imports/template.xlsx", headers=h)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert wb["students"].data_validations.dataValidation, "drop-down lists present"
    codes = [c.value for c in wb["lists"]["F"][1:] if c.value]
    assert "BTECH-CSE" in codes and "MCA" not in codes  # own programme codes only
    _ = Workbook


# ---------------------------------------------------------------- upload hardening (parse_upload)

import time  # noqa: E402
import zipfile  # noqa: E402

import pytest  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
from openpyxl.worksheet._read_only import ReadOnlyWorksheet  # noqa: E402

from app.config import settings  # noqa: E402
from app.modules import intake  # noqa: E402
from app.modules.normalize import TEMPLATE_COLUMNS  # noqa: E402
from app.security import ApiError  # noqa: E402

ROW = ["911100000001", "KIT26X", "Test Person", "2008-05-12", "F", "INDIAN", "", "BTECH-CSE", "2026-27",
       "REGULAR", "2026-08-03", "", "", ""]


def _xlsx(n_rows: int, blank_styled_to: int = 0) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(TEMPLATE_COLUMNS)
    for i in range(n_rows):
        ws.append([*ROW[:1], f"R{i}", *ROW[2:]])
    fill = PatternFill("solid", fgColor="FFFF00")
    for r in range(n_rows + 2, blank_styled_to + 1):
        ws.cell(row=r, column=1).fill = fill
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(n_rows: int) -> bytes:
    lines = [",".join(TEMPLATE_COLUMNS)] + [",".join([ROW[0], f"R{i}", *ROW[2:]]) for i in range(n_rows)]
    return ("\n".join(lines) + "\n").encode()


def test_xlsx_zip_bomb_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("xl/worksheets/sheet1.xml", b" " * (60 * 1024 * 1024))
    data = buf.getvalue()
    assert len(data) < settings.max_upload_bytes
    t = time.monotonic()
    with pytest.raises(ApiError) as e:
        intake.parse_upload("b.xlsx", data)
    assert e.value.status == 413 and "too large when unpacked" in e.value.message
    assert time.monotonic() - t < 2


def test_xlsx_doctype_rejected():
    src = zipfile.ZipFile(io.BytesIO(_xlsx(1)))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                body = b'<!DOCTYPE x [<!ENTITY a "aaaa">]>' + body
            z.writestr(item, body)
    with pytest.raises(ApiError) as e:
        intake.parse_upload("d.xlsx", buf.getvalue())
    assert e.value.status == 422 and "cannot be read safely" in e.value.message


def test_row_cap_boundary():
    assert len(intake.parse_upload("ok.csv", _csv(5000))) == 5000
    with pytest.raises(ApiError) as e:
        intake.parse_upload("over.csv", _csv(5001))
    assert e.value.status == 413 and "5000" in e.value.message


def test_xlsx_blank_trailing_rows_accepted():
    assert len(intake.parse_upload("blank.xlsx", _xlsx(3, blank_styled_to=3000))) == 3


def test_xlsx_rows_streamed_not_materialised(monkeypatch):
    data = _xlsx(50)
    yielded = 0
    original = ReadOnlyWorksheet.iter_rows

    def counting(self, *a, **kw):
        nonlocal yielded
        for row in original(self, *a, **kw):
            yielded += 1
            yield row

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", counting)
    monkeypatch.setattr(settings, "max_upload_rows", 10)
    with pytest.raises(ApiError) as e:
        intake.parse_upload("s.xlsx", data)
    assert e.value.status == 413
    assert yielded < 15
