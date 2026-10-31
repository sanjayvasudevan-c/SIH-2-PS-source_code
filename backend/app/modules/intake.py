"""Bulk import: parse -> preview (validate + projection) -> correct / exclude / bulk-fix -> commit.

Commit takes a job lock with a conditional UPDATE, re-validates every row, and commits each row
in its own savepoint so one failure never loses the others (partial commit, then retry).
"""

from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from collections import Counter
from datetime import timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import and_, or_, select, update
from sqlalchemy.orm import Session

from .. import audit
from ..config import settings
from ..csvout import to_csv
from ..db import restore_hooks, snapshot_hooks, utcnow
from ..models import ImportJob, ImportRow, Institution, Programme
from ..security import Actor, ApiError
from . import engine, issues, ledger
from .normalize import (
    GENDERS,
    PENDING_REASONS,
    REQUIRED_COLUMNS,
    SEAT_CATEGORIES,
    SOCIAL_CATEGORIES,
    TEMPLATE_COLUMNS,
    canonical_header,
    clean_cell,
    finding,
)
from .rules import load_config

EDITABLE = set(TEMPLATE_COLUMNS)
STALE_LOCK = timedelta(minutes=10)


# ---------------------------------------------------------------- parsing


MAX_XLSX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_XLSX_RATIO = 100


def _check_xlsx_container(content: bytes) -> None:
    """Reject zip bombs and XML with DOCTYPE/entity declarations before openpyxl unpacks anything."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(content))
        members = zf.infolist()
    except zipfile.BadZipFile as exc:
        raise ApiError(415, "This file is not a real Excel (.xlsx) workbook.") from exc
    too_big = sum(m.file_size for m in members) > MAX_XLSX_UNCOMPRESSED or any(
        m.file_size / max(m.compress_size, 1) > MAX_XLSX_RATIO for m in members)
    if too_big:
        raise ApiError(413, "This workbook is too large when unpacked. Save it again from Excel, or use the CSV template.")
    with zf:
        for m in members:
            if m.filename.startswith("xl/") and m.filename.endswith(".xml"):
                with zf.open(m) as f:
                    if b"<!DOCTYPE" in f.read(2048):
                        raise ApiError(422, "This workbook cannot be read safely. Save it again from Excel and retry.")


def parse_upload(filename: str, content: bytes) -> list[dict]:
    name = (filename or "").lower()
    if not (name.endswith(".csv") or name.endswith(".xlsx")):
        raise ApiError(415, "Upload a .csv or .xlsx file (use the template).")
    if len(content) > settings.max_upload_bytes:
        raise ApiError(413, "The file is larger than 2 MB. Split it into smaller files.")
    wb = None
    if name.endswith(".xlsx"):
        if not content.startswith(b"PK"):
            raise ApiError(415, "This file is not a real Excel (.xlsx) workbook.")
        _check_xlsx_container(content)
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ApiError(422, "The Excel file could not be read. Save it again from Excel and retry.") from exc
        it = wb.worksheets[0].iter_rows(values_only=True)
    else:
        if b"\x00" in content:
            raise ApiError(415, "The CSV file contains binary data. Save it as plain CSV (UTF-8).")
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ApiError(422, "The CSV is not UTF-8 encoded. In Excel use 'CSV UTF-8' when saving.") from exc
        it = csv.reader(io.StringIO(text))
    try:
        header = next(it, None) or ()
        cols = [canonical_header(h) for h in header]
        missing = [c for c in REQUIRED_COLUMNS if c not in cols]
        if missing:
            raise ApiError(422, "The header row is missing required columns: " + ", ".join(missing)
                           + ". Download the template to see the expected columns.", missing_columns=missing)
        rows = []
        for r in it:  # streamed: stop as soon as the row cap is passed
            if all(clean_cell(v) == "" for v in r):
                continue
            if len(rows) >= settings.max_upload_rows:
                raise ApiError(413, f"The file has more than {settings.max_upload_rows} rows; "
                               f"the limit is {settings.max_upload_rows} per upload. Split it into smaller files.")
            rows.append({cols[i]: clean_cell(r[i]) if i < len(r) else "" for i in range(len(cols)) if cols[i]})
    finally:
        if wb is not None:
            wb.close()
    if not rows:
        raise ApiError(422, "The file has a header but no data rows.")
    return rows


# ---------------------------------------------------------------- jobs


def get_job(db: Session, job_id: str, actor: Actor) -> ImportJob:
    job = db.get(ImportJob, job_id)
    if job is None or (actor.is_institution and job.institution_id != actor.institution_id):
        raise ApiError(404, "Not found.")
    return job


def create_job(db: Session, actor: Actor, institution: Institution, academic_year: str, filename: str,
               content: bytes, rows: list[dict] | None = None) -> ImportJob:
    rows = rows if rows is not None else parse_upload(filename, content)
    sha = hashlib.sha256(content).hexdigest()
    prior = db.execute(
        select(ImportJob).where(ImportJob.institution_id == institution.id, ImportJob.file_sha256 == sha)
        .order_by(ImportJob.created_at)
    ).scalars().first()
    job = ImportJob(institution_id=institution.id, academic_year=academic_year,
                    uploaded_by=actor.id if actor.kind == "USER" else None, filename=filename[:255],
                    file_sha256=sha, status="PREVIEW", total_rows=len(rows), summary={})
    db.add(job)
    db.flush()
    for i, r in enumerate(rows):
        db.add(ImportRow(job_id=job.id, row_no=i, data=r, original=dict(r), status="VALID", findings=[]))
    db.flush()
    file_findings = []
    cfg = load_config(db)
    if prior is not None and cfg["F004_REPEATED_FILE"].enabled:
        f = finding("F004_REPEATED_FILE", cfg["F004_REPEATED_FILE"].severity, None,
                    f"This exact file was uploaded before ({prior.filename}, {prior.created_at:%d %b %Y %H:%M} UTC).",
                    {"previous_job_id": prior.id})
        file_findings.append(f)
    job.summary = {"file_findings": file_findings}
    revalidate(db, job)
    audit.record(db, actor, "import.upload", "import_job", job.id,
                 detail={"filename": filename, "rows": len(rows), "sha256": sha, "counts": job.summary["counts"]},
                 institution_id=institution.id)
    return job


def _active_rows(db: Session, job: ImportJob) -> list[ImportRow]:
    return list(db.execute(select(ImportRow).where(ImportRow.job_id == job.id).order_by(ImportRow.row_no)).scalars())


def revalidate(db: Session, job: ImportJob) -> list[tuple[ImportRow, engine.RowPlan]]:
    institution = db.get(Institution, job.institution_id)
    rows = _active_rows(db, job)
    pending = [r for r in rows if r.status not in ("COMMITTED", "EXCLUDED")]
    plans = engine.validate_rows(db, institution, job.academic_year, [(r.row_no, r.data) for r in pending])
    pairs = list(zip(pending, plans, strict=True))
    for r, p in pairs:
        r.data = p.data
        r.findings = p.findings
        r.status = p.status
    db.flush()
    _summarise(db, job, rows, pairs)
    return pairs


def _summarise(db: Session, job: ImportJob, rows: list[ImportRow], pairs) -> None:
    counts = Counter(r.status for r in rows)
    rule_counts = Counter(f["rule_id"] for r in rows if r.status not in ("COMMITTED", "EXCLUDED") for f in r.findings)
    summary = dict(job.summary or {})
    summary["counts"] = {s: counts.get(s, 0) for s in ("VALID", "WARNING", "REVIEW", "ERROR", "COMMITTED", "EXCLUDED")}
    summary["committable"] = sum(counts.get(s, 0) for s in ("VALID", "WARNING", "REVIEW"))
    summary["rules"] = dict(rule_counts)
    summary["projection"] = projection(db, job, pairs)
    job.summary = summary


def projection(db: Session, job: ImportJob, pairs) -> list[dict]:
    """Current vs after-commit seat position per programme touched by committable rows."""
    from ..models import Enrolment

    deltas: dict[str, Counter] = {}
    for _row, p in pairs:
        if p.status == "ERROR" or not p.programme_id:
            continue
        year = p.parsed["academic_year"] or job.academic_year
        if year != job.academic_year:
            continue
        deltas.setdefault(p.programme_id, Counter())[p.parsed["seat_category"]] += 1
        if p.enrolment_id:
            old = db.get(Enrolment, p.enrolment_id)
            if old and ledger.occupies_seat(old.status) and old.academic_year == job.academic_year:
                deltas.setdefault(old.programme_id, Counter())[old.seat_category] -= 1
    out = []
    for pid, delta in deltas.items():
        prog = db.get(Programme, pid)
        intake = ledger.get_intake(db, pid, job.academic_year)
        caps = ledger.caps_for(db, intake)
        counts = ledger.counts_for(db, pid, job.academic_year)
        after_counts = Counter(counts)
        after_counts.update(delta)
        n = intake.approved_intake if intake else 0
        cur = ledger.position_from_counts(n, caps, counts)
        aft = ledger.position_from_counts(n, caps, dict(after_counts))
        new_excess = aft["excess_over_intake"] > cur["excess_over_intake"] or any(
            aft["supernumerary"][c]["excess"] > cur["supernumerary"][c]["excess"] for c in aft["supernumerary"]
        )
        out.append({
            "programme_id": pid,
            "programme_code": prog.code,
            "programme_name": prog.name,
            "current": cur,
            "after": aft,
            "change_within_intake": aft["within_intake_enrolled"] - cur["within_intake_enrolled"],
            "exceeds": aft["any_excess"],
            "message": (
                "Committing will exceed approved capacity. Records are still accepted and AICTE is flagged."
                if new_excess or aft["any_excess"] else "Committing stays within approved capacity."
            ),
        })
    out.sort(key=lambda x: x["programme_code"])
    return out


def _editable_row(db: Session, job: ImportJob, row_no: int) -> ImportRow:
    if job.status in ("COMMITTED", "DISCARDED", "COMMITTING"):
        raise ApiError(409, f"This upload is {job.status.lower()}; rows can no longer be changed.")
    row = db.execute(select(ImportRow).where(ImportRow.job_id == job.id, ImportRow.row_no == row_no)).scalar_one_or_none()
    if row is None:
        raise ApiError(404, "Row not found.")
    if row.status == "COMMITTED":
        raise ApiError(409, "This row is already committed. Correct the student record instead.",
                       row_status=row.status, enrolment_id=row.enrolment_id)
    return row


def update_row(db: Session, job: ImportJob, row_no: int, changes: dict, actor: Actor) -> ImportRow:
    row = _editable_row(db, job, row_no)
    bad = [k for k in changes if k not in EDITABLE]
    if bad:
        raise ApiError(422, "Unknown field(s): " + ", ".join(bad))
    data = dict(row.data)
    for k, v in changes.items():
        data[k] = clean_cell(v)
    row.data = data
    row.corrected = True
    revalidate(db, job)
    audit.record(db, actor, "import.correct", "import_row", row.id,
                 detail={"job_id": job.id, "row_no": row_no, "fields": sorted(changes)})
    return row


def exclude_row(db: Session, job: ImportJob, row_no: int, excluded: bool, actor: Actor) -> ImportRow:
    row = _editable_row(db, job, row_no)
    if excluded:
        row.status = "EXCLUDED"
    elif row.status == "EXCLUDED":
        row.status = "VALID"  # recomputed below
    db.flush()
    revalidate(db, job)
    audit.record(db, actor, "import.exclude" if excluded else "import.restore", "import_row", row.id,
                 detail={"job_id": job.id, "row_no": row_no})
    return row


def bulk_fix(db: Session, job: ImportJob, field: str, from_value: str, to_value: str, actor: Actor) -> int:
    if job.status in ("COMMITTED", "DISCARDED", "COMMITTING"):
        raise ApiError(409, f"This upload is {job.status.lower()}; rows can no longer be changed.")
    if field not in EDITABLE:
        raise ApiError(422, f"Unknown field {field}.")
    n = 0
    for r in _active_rows(db, job):
        if r.status in ("COMMITTED",):
            continue
        if (r.data.get(field) or "") == (from_value or ""):
            r.data = dict(r.data, **{field: clean_cell(to_value)})
            r.corrected = True
            n += 1
    db.flush()
    revalidate(db, job)
    audit.record(db, actor, "import.bulk_fix", "import_job", job.id,
                 detail={"field": field, "rows": n})
    return n


def commit_job(db: Session, job: ImportJob, actor: Actor) -> dict:
    if job.status in ("COMMITTED", "DISCARDED"):
        return {"job_id": job.id, "status": job.status, "committed_now": 0,
                "message": f"Nothing to do: this upload is already {job.status.lower()}."}
    previous = job.status
    now = utcnow()
    res = db.execute(
        update(ImportJob)
        .where(
            ImportJob.id == job.id,
            or_(
                ImportJob.status.in_(("PREVIEW", "PARTIALLY_COMMITTED")),
                and_(ImportJob.status == "COMMITTING", ImportJob.updated_at < now - STALE_LOCK),
            ),
        )
        .values(status="COMMITTING", updated_at=now)
        .execution_options(synchronize_session=False)
    )
    if res.rowcount != 1:
        db.rollback()
        raise ApiError(409, "This upload is already being committed. Wait a moment and refresh.")
    db.commit()
    db.refresh(job)
    institution = db.get(Institution, job.institution_id)
    try:
        pairs = revalidate(db, job)
        committed_now = 0
        failed = 0
        affected = set()
        for row, plan in pairs:
            if plan.status == "ERROR":
                continue
            snap = snapshot_hooks(db)
            sp = db.begin_nested()
            try:
                en = engine.commit_values(db, plan, institution=institution, academic_year=job.academic_year,
                                          actor=actor, source="IMPORT", import_job_id=job.id,
                                          reason=f"Committed from upload {job.filename} row {row.row_no + 2}")
                sp.commit()
            except Exception as exc:  # noqa: BLE001
                sp.rollback()
                restore_hooks(db, snap)
                failed += 1
                row.status = "ERROR"
                row.findings = [*row.findings, finding("COMMIT_FAILED", "ERROR", None,
                                                        "This row could not be saved; it was left uncommitted. Retry the commit.",
                                                        {"error": type(exc).__name__})]
                continue
            row.status = "COMMITTED"
            row.enrolment_id = en.id
            committed_now += 1
            affected.add((en.programme_id, en.academic_year))
        for pid, year in sorted(affected):
            issues.run_capacity_checks(db, pid, year, actor)
        rows = _active_rows(db, job)
        remaining = [r for r in rows if r.status not in ("COMMITTED", "EXCLUDED")]
        job.committed_rows += committed_now
        if not remaining:
            job.status = "COMMITTED"
        elif job.committed_rows > 0:
            job.status = "PARTIALLY_COMMITTED"
        else:
            job.status = previous if previous != "COMMITTING" else "PREVIEW"
        _summarise(db, job, rows, [])
        audit.record(db, actor, "import.commit", "import_job", job.id,
                     detail={"committed_now": committed_now, "failed": failed, "remaining": len(remaining),
                             "status": job.status}, institution_id=job.institution_id)
        db.commit()
    except Exception:
        db.rollback()
        db.execute(update(ImportJob).where(ImportJob.id == job.id).values(status=previous if previous != "COMMITTING" else "PREVIEW"))
        db.commit()
        raise
    return {
        "job_id": job.id,
        "status": job.status,
        "committed_now": committed_now,
        "committed_total": job.committed_rows,
        "remaining": len(remaining),
        "message": f"Committed {committed_now} row(s)."
        + (f" {len(remaining)} row(s) still need attention." if remaining else ""),
    }


def discard_job(db: Session, job: ImportJob, actor: Actor) -> None:
    if job.status in ("COMMITTED", "COMMITTING"):
        raise ApiError(409, "A committed upload cannot be discarded; correct individual records instead.")
    job.status = "DISCARDED"
    audit.record(db, actor, "import.discard", "import_job", job.id)


def report_csv(db: Session, job: ImportJob) -> str:
    out = []
    for r in _active_rows(db, job):
        base = [r.row_no + 2, r.status, r.data.get("institution_student_id"), r.data.get("full_name")]
        if not r.findings:
            out.append([*base, "", "", "", "", ""])
        for f in r.findings:
            sugg = (f.get("suggestion") or {}).get("value") if f.get("suggestion") else ""
            out.append([*base, f["rule_id"], f["severity"], f.get("field") or "", f["message"], sugg or ""])
    return to_csv(["spreadsheet_row", "row_status", "institution_student_id", "full_name", "rule_id", "severity",
                   "field", "message", "suggested_value"], out)


# ---------------------------------------------------------------- templates


def template_csv() -> str:
    example = ["", "ROLL001", "Example Student", "2007-05-21", "F", "INDIAN", "India", "PROGRAMME-CODE", "2026-27",
               "REGULAR", "2026-08-01", "", "student@example.org", "APPLIED_AWAITING_ID"]
    return to_csv(TEMPLATE_COLUMNS, [example])


def template_xlsx(programme_codes: list[str], academic_year: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "students"
    ws.append(TEMPLATE_COLUMNS)
    lists = wb.create_sheet("lists")
    columns = {
        "gender": sorted(set(GENDERS.values())),
        "nationality_status": ["INDIAN", "NRI", "OCI", "FN"],
        "seat_category": list(SEAT_CATEGORIES),
        "apaar_pending_reason": list(PENDING_REASONS),
        "social_category": list(SOCIAL_CATEGORIES),
        "programme_code": programme_codes or ["(none)"],
        "academic_year": [academic_year],
    }
    for ci, (col, values) in enumerate(columns.items(), start=1):
        letter = lists.cell(row=1, column=ci).column_letter
        lists.cell(row=1, column=ci, value=col)
        for ri, v in enumerate(values, start=2):
            lists.cell(row=ri, column=ci, value=v)
        target = ws.cell(row=1, column=TEMPLATE_COLUMNS.index(col) + 1).column_letter
        dv = DataValidation(type="list", formula1=f"=lists!${letter}$2:${letter}${len(values) + 1}", allow_blank=True)
        dv.error = f"Choose a value from the {col} list."
        dv.errorTitle = "Not in list"
        ws.add_data_validation(dv)
        dv.add(f"{target}2:{target}5001")
    for c in ("date_of_birth", "admission_date"):
        letter = ws.cell(row=1, column=TEMPLATE_COLUMNS.index(c) + 1).column_letter
        for row in range(2, 5002):
            ws[f"{letter}{row}"].number_format = "yyyy-mm-dd"
    apaar_letter = ws.cell(row=1, column=1).column_letter
    for row in range(2, 5002):
        ws[f"{apaar_letter}{row}"].number_format = "@"
    for i, _ in enumerate(TEMPLATE_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 20
    lists.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
